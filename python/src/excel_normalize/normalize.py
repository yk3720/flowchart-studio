from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .constants import DEFAULT_LAYOUT, FLOW_SCHEMA
from .kosei import parse_kosei_sheet
from .tables import extract_unit_sheet_tables, list_unit_sheets
from .validate import NormalizeError, NormalizeResult, validate_bundle


def _build_units(kosei) -> list[dict]:
    units: list[dict] = []
    unit_sort: dict[str, int] = {}
    module_sort: dict[tuple[str, str], int] = {}
    module_memos: dict[tuple[str, str], str] = {}
    unit_memos: dict[str, str] = {}

    for row in kosei.rows:
        if row.unit_label not in unit_sort:
            unit_sort[row.unit_label] = len(unit_sort)
        module_sort[(row.unit_label, row.module_label)] = row.sort_index
        if row.module_memo:
            module_memos[(row.unit_label, row.module_label)] = row.module_memo

    if kosei.unit_bands:
        for band in kosei.unit_bands.values():
            if band.memo:
                unit_memos[band.label] = band.memo

    for unit_label in kosei.unit_labels:
        modules = []
        for mod in kosei.modules_for_unit(unit_label):
            entry: dict = {
                "label": mod,
                "sort_order": module_sort[(unit_label, mod)],
            }
            memo = module_memos.get((unit_label, mod))
            if memo:
                entry["memo"] = memo
            modules.append(entry)
        unit_entry: dict = {
            "label": unit_label,
            "sort_order": unit_sort[unit_label],
            "modules": modules,
        }
        unit_memo = unit_memos.get(unit_label)
        if unit_memo:
            unit_entry["memo"] = unit_memo
        units.append(unit_entry)
    return units


def _build_payload(module_label: str, rows: list[list[str]]) -> dict:
    return {
        "version": 1,
        "schema": FLOW_SCHEMA,
        "title": module_label,
        "layout": dict(DEFAULT_LAYOUT),
        "table": rows,
        "createdAt": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z"),
    }


def normalize_workbook(workbook_path: Path) -> NormalizeResult:
    wb = load_workbook(workbook_path, data_only=True)
    kosei = parse_kosei_sheet(wb)
    unit_sheets = list_unit_sheets(wb, kosei)

    flows: list[dict] = []
    for unit_label, ws in unit_sheets.items():
        expected_modules = kosei.modules_for_unit(unit_label)
        blocks = extract_unit_sheet_tables(ws, unit_label, expected_modules, kosei)
        block_by_module = {b.module_label: b for b in blocks}

        for module_label in expected_modules:
            block = block_by_module.get(module_label)
            if block is None:
                continue
            flows.append(
                {
                    "unit_label": unit_label,
                    "module_label": module_label,
                    "title": module_label,
                    "payload": _build_payload(module_label, block.rows),
                }
            )

    errors, warnings = validate_bundle(kosei, flows)
    if errors:
        raise NormalizeError(errors)

    bundle = {
        "internal_code": kosei.internal_code,
        "display_name": kosei.display_name,
        "units": _build_units(kosei),
        "flows": flows,
    }
    if kosei.device_memo:
        bundle["memo"] = kosei.device_memo
    return NormalizeResult(bundle=bundle, warnings=warnings)
